from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(r"D:\work\shopkeeper_brain-1")
OUT = ROOT / "doc"
OUT.mkdir(parents=True, exist_ok=True)

def extract_docx(path):
    doc = Document(path)
    lines = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            lines.append(text)
    return lines

lc = extract_docx(ROOT / "aaa" / "LangChainV1.0.2.docx")
lg = extract_docx(ROOT / "aaa" / "LangGraphV1.0.3.docx")

slides = [
    ("掌柜智库：听书知识库 RAG 系统", ["10分钟答辩 · 项目功能、结构、流程与差异"], "开场：本项目面向听书内容管理与问答，不只是把文档接入大模型，而是把音频、书籍元数据、知识图谱、检索和评估串成可验证的业务闭环。"),
    ("1｜有什么用：产品与领域", ["面向听书平台的内容资产管理", "MP3 自动转写，沉淀为可检索知识", "回答书籍推荐、详情、内容检索、听书笔记", "来源可追溯，库外问题明确拒答"], "先讲价值。听书平台的核心问题不是缺少文本，而是音频内容难检索、书籍信息分散、回答难以溯源。本项目将音频和文档统一加工成知识单元，服务听众、内容运营和编辑三类角色。"),
    ("2｜项目结构：四层 + 两条工作流", ["API/Service：导入、查询、评估三类服务", "Processor：ImportGraph 与 QueryGraph", "Storage：Milvus、Neo4j、MongoDB、MinIO", "Eval：数据集、指标、报告、刷新与 UI"], "结构上分四层。最重要的是两条 LangGraph 工作流：导入图负责把原始材料变成知识，查询图负责把问题变成可追溯答案；评估模块独立出来，避免只凭感觉调参。"),
    ("3｜项目执行顺序：从音频到答案", ["导入：MP3/PDF/Markdown/HTML → 转换 → 切片", "增强：BGE-M3 向量、书籍字段、知识图谱", "查询：问题改写 → 多路检索 → RRF → 重排", "输出：答案、来源、会话记录、SSE 流式返回"], "按执行顺序看，先导入再查询。音频经过转写和切片后，与其他材料进入同一知识库；查询时结合向量、HyDE、图谱和必要的网络路径，最后经过融合与重排，输出带来源的答案。"),
    ("4｜模型简介：重点看语音转文字", ["LLM：负责字段识别、问题改写、答案生成", "BGE-M3：生成稠密 + 稀疏向量", "BGE Reranker：对候选片段二次排序", "Whisper：MP3 → 中文文本，是本项目新增重点"], "模型只做快速交代。LLM、BGE-M3 和重排模型是常规 RAG 组件。重点是 Whisper：它把原本只能播放的 MP3 转成带时间和内容的文本，再走与文档相同的切片、向量化和知识图谱流程，音频因此获得检索能力。"),
    ("5｜重点代码：导入链路", ["mp3_to_md_node.py：音频转写与中间文件", "book_info_recognition_node.py：书名、作者、类型、条目等字段", "document_split_node.py：统一切片", "bge_embedding_chunks_node.py + import_milvus_node.py：向量入库", "kg_graph_node.py：实体关系入 Neo4j"], "不贴代码，只讲职责。导入链路的关键是统一：音频不是单独的一套系统，而是先转成 Markdown，再复用切片、字段识别、向量和图谱节点。这样后续检索不需要区分来源格式。"),
    ("6｜重点代码：查询、评估与本项目差异", ["query_process：HyDE、向量/图谱/网络多路检索、RRF、重排", "eval/：金标集、越界召回、库外拒答、P/R/F1、生成层指标", "音频元数据：audio_duration、entry_name、source_file_name", "相对 shopkeeper_brain：新增听书场景、音频导入、评估中心与报告体系"], "这里重点对比基线项目。基础导入、查询和存储能力相似，略过。新增价值在四点：MP3 转写；听书专属字段和场景；可重复运行的评估中心；库外拒答和来源质量约束。"),
    ("7｜测试与评估：用数据证明效果", ["golden_set：检索 P/R/F1、Macro/Micro-F1、MRR", "cross_recall：检测问 A 书召回 B 书", "negative_set：检测库外拒答与澄清式拒答", "报告 JSON + Markdown，支持重试、刷新、性能耗时"], "项目不是导入后凭感觉演示。金标集看召回质量，越界集看串书，库外集看是否胡答；评估还记录生成层引用准确率、忠实度和回答相关性，并通过导入清单保证刷新后可复现。"),
    ("8｜总结：从可用 RAG 到可评估听书知识库", ["业务闭环：音频资产 → 知识 → 问答 → 评估", "技术重点：LangGraph 编排 + Whisper 转写 + 混合检索", "工程重点：来源字段、增量导入、库外拒答、可复现报告", "答辩结束：谢谢"], "总结三句话：第一，这是听书领域的知识库产品；第二，核心增量是音频转写和听书场景字段；第三，项目把评估和拒答纳入工程闭环，能够说明系统答得怎样，而不只是展示一次成功回答。"),
]

def add_textbox(slide, x, y, w, h, text, size=24, color=(40,38,34), bold=False):
    box = slide.shapes.add_textbox(PInches(x), PInches(y), PInches(w), PInches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text; p.font.size = PPt(size); p.font.bold = bold; p.font.color.rgb = RGBColor(*color)
    return box

prs = Presentation(); prs.slide_width = PInches(13.333); prs.slide_height = PInches(7.5)
for i, (title, bullets, notes) in enumerate(slides):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = RGBColor(248, 246, 241)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, PInches(0.22)); band.fill.solid(); band.fill.fore_color.rgb = RGBColor(194, 126, 51); band.line.fill.background()
    add_textbox(slide, 0.7, 0.65, 11.8, 0.7, title, 28, (42,38,34), True)
    for j, bullet in enumerate(bullets):
        add_textbox(slide, 1.0, 1.75 + j * 0.78, 11.3, 0.52, "• " + bullet, 20, (74,66,57))
    add_textbox(slide, 0.7, 6.9, 11.8, 0.25, f"Shopkeeper Brain · {i+1}/{len(slides)} · 约 {10/len(slides):.1f} 分钟", 10, (120,108,95))
    slide.notes_slide.notes_text_frame.text = "讲解备注：" + notes
prs.save(OUT / "掌柜智库_10分钟答辩.pptx")

def setup_doc(title):
    d = Document(); sec = d.sections[0]; sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run(title); r.bold = True; r.font.size = Pt(20)
    return d

script = setup_doc("掌柜智库 10 分钟答辩演讲稿与 PPT 备注")
script.add_paragraph("建议总时长：8.5-10分钟。每页控制在约1分钟，模型介绍和重复基础架构快速带过，重点突出音频转写、听书场景和评估闭环。")
for i, (title, _, notes) in enumerate(slides, 1):
    script.add_heading(f"第{i}页｜{title}", level=1)
    script.add_paragraph(notes)
    script.add_paragraph("备注：本页不贴代码，以模块职责、数据流和项目价值为主。")
script.add_heading("答辩可能追问", level=1)
for q in ["为什么需要 Whisper？答：音频是听书平台的主要内容形态，转写后才能统一检索、切片、向量化和溯源。", "为什么要做评估模块？答：RAG 的改动需要可量化比较，金标、越界召回和库外拒答分别覆盖正确性、串书和胡答风险。", "与原项目最主要区别？答：不是基础框架重写，而是增加听书领域能力和质量闭环：音频导入、字段体系、业务场景、评估中心。"]:
    script.add_paragraph(q, style="List Bullet")
script.save(OUT / "掌柜智库_答辩演讲稿与备注.docx")

report = setup_doc("掌柜智库项目报告")
report.add_paragraph("版本：shopkeeper_brain-1｜用途：项目答辩配套报告｜日期：2026-09-05")
sections = [
    ("1. 需求分析", "面向听书平台建立可管理、可检索、可追溯的知识库。需求包括多格式内容导入，重点支持 MP3 转写；书籍推荐、详情、内容检索、听书笔记四类问答；多角色使用；来源展示；库外拒答；以及可重复的离线评估。"),
    ("2. 产品与领域设计", "产品对象是听众、内容运营和内容编辑。领域对象包括书籍、作者、条目、内容类型、类别标签、有声书时长、来源文件和知识切片。系统把音频内容转化为与文档一致的知识单元，支持书籍级和条目级定位。"),
    ("3. 总体架构", "系统采用 FastAPI + LangGraph。API 层承接导入、查询和评估；工作流层分为 ImportGraph 与 QueryGraph；工具层提供 LLM、BGE-M3、重排、SSE 和任务管理；存储层使用 Milvus、Neo4j、MongoDB、MinIO。"),
    ("4. 导入设计", "导入图根据扩展名路由。PDF/HTML/Markdown/TXT 进入文本处理，MP3 进入 mp3_to_md_node 完成语音转文字，随后统一经过内容识别、文档切片、BGE-M3 向量化、Milvus 入库和 Neo4j 图谱构建。增量导入复用已有集合，评估通过 manifest 管理材料。"),
    ("5. 查询设计", "查询图先做问题改写和书名确认，再并行执行向量检索、HyDE 检索、知识图谱检索及必要的网络检索；通过 RRF 融合和 BGE 重排得到最终上下文；回答节点依据角色和场景提示词生成答案，并输出来源区块和 SSE 流式结果。"),
    ("6. 模型与技术选型", "LLM 用于结构化识别、查询改写和答案生成；BGE-M3 提供稠密和稀疏向量；BGE Reranker 做候选重排；Whisper 是本项目重点，用于把 MP3 转为中文文本。LangChain 负责模型和检索组件抽象，LangGraph 负责有状态流程编排。"),
    ("7. 代码实现", "核心代码按职责拆分：api/services 管理接口和后台任务；processor/import_process 实现导入节点；processor/query_process 实现查询节点；prompts 管理场景提示词；utils 封装存储和模型客户端；front 提供导入、聊天和评估页面；eval 提供数据集、指标、报告和刷新。"),
    ("8. 相对 shopkeeper_brain 的差异", "基础 FastAPI、Milvus、Neo4j、MongoDB、MinIO 和 RAG 思路相似，因此不重复展开。本项目重点新增或强化：MP3 转写导入；听书专属元数据；书籍推荐/详情/听书笔记场景；评估中心；金标、越界召回、库外拒答；来源字段透传；增量导入和评估材料隔离。"),
    ("9. 测试与评估", "测试覆盖服务连接、导入节点、向量检索、图谱和 RAG。评估使用 golden_set 计算 P/R/F1、Micro/Macro-F1、MRR；cross_recall 检查串书；negative_set 检查库外拒答和澄清式拒答；生成层补充引用准确率、关键词准确率、字符 F1、忠实度和答案相关性。"),
    ("10. 结论与后续", "项目形成音频资产到知识、问答和评估的闭环。当前重点风险是外部模型/数据库依赖和真实数据端到端验证；后续可增加统一字段常量、更多音频质量指标、报告趋势对比、API 自动化测试和来源链接的完整展示。"),
]
for h, body in sections:
    report.add_heading(h, level=1); report.add_paragraph(body)
    report.add_page_break()
report.add_heading("参考资料", level=1)
report.add_paragraph("本报告吸收并结合 LangChainV1.0.2.docx 与 LangGraphV1.0.3.docx 中关于组件抽象、模型调用链和图工作流编排的说明；两份文档位于项目 aaa 目录。项目目录中的 LangGraph 参考文件实际版本为 V1.0.3。")
report.add_heading("附录：参考文档提要", level=1)
report.add_paragraph("LangChain 文档要点：通过统一接口连接模型、提示词、解析器、检索器和工具，适合构建 RAG 组件链。LangGraph 文档要点：用节点、边、状态和条件路由表达有状态、可恢复、可观测的流程。本项目分别对应查询/导入节点的组件化和 ImportGraph/QueryGraph 的流程编排。")
report.save(OUT / "掌柜智库_项目报告_10页.docx")

print("generated", OUT)
